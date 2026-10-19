#!/usr/bin/env python3
"""Render the Phase 0 kick-off pack and the decision sheet as one branded Word document.

Companion to `export_storyboard_docx.py` (it reuses that module's Markdown renderer, branded
header/footer and document-control table). Output: `00_program/kickoff-pack.docx`, document id
`MC-KO-01`, sent with the kick-off invite alongside the nine storyboard renderings.

Usage:
    python3 elearning-course/scorm-plans/tools/export_kickoff_docx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_storyboard_docx import (  # noqa: E402
    DATE, LOGO, brand_footer, brand_header, doc_control_table, render_markdown_body, rgb,
)

ROOT = Path(__file__).resolve().parents[3]
PLANS = ROOT / "elearning-course" / "scorm-plans"
PROGRAM = PLANS / "00_program"
DOC_ID = "MC-KO-01 v0.3"
NAVY = "003867"


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    normal = doc.styles["Normal"]
    normal.font.name = "Montserrat"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Montserrat")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11), ("Heading 4", 10)):
        style = doc.styles[name]
        style.font.name = "Montserrat"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Montserrat")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(NAVY)


def main() -> None:
    doc = Document()
    setup(doc)
    brand_header(doc.sections[0], None, DOC_ID)
    brand_footer(doc.sections[0], f"Phase 0 kick-off working document · 2026-10-19 · {DOC_ID}")

    band = doc.add_table(rows=1, cols=1)
    from export_storyboard_docx import set_cell_shading
    set_cell_shading(band.rows[0].cells[0], NAVY)
    band.rows[0].cells[0].paragraphs[0].add_run("")
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.add_run().add_picture(str(LOGO), width=Inches(1.72))
    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(3)
    title.add_run("My Clinic — Phase 0 kick-off pack")
    strap = doc.add_paragraph()
    strap.add_run("Kick-off working document · ").bold = True
    strap.add_run("Agenda, attendees, pre-reads and the complete open-decision list with owners and "
                  "due dates. Recommended answers are drafts for the named owner to confirm in "
                  "writing; nothing here is approved course content.")
    strap.paragraph_format.space_after = Pt(10)
    doc_control_table(doc, [
        ("Document", "MC-KO-01 · Phase 0 kick-off pack + decision sheet"),
        ("Revision", "v0.3 · 2026-10-19 (researched working assumptions added)"),
        ("Proposed kick-off", "Monday 2026-10-19 (week 1 of the proposal timeline)"),
        ("Status", "Working document — decisions are recorded in the Phase 0 files, not here"),
        ("Sources", "00_program/kickoff-pack.md + 00_program/kickoff-assumed-answers.md (generated)"),
        ("Companion files", "client-review pack (9 Word docs) · decision-sheet.csv · "
                            "storyboard-review-register.csv · proposal Rev01"),
    ], NAVY)
    doc.add_page_break()

    render_markdown_body(doc, (PROGRAM / "kickoff-pack.md").read_text(encoding="utf-8").splitlines(), NAVY)
    doc.add_page_break()
    heading = doc.add_heading("Attachment — open decisions with researched working assumptions", level=1)
    heading.paragraph_format.page_break_before = False
    # the full tracker (owner, due date, write-back file) is decision-sheet.csv; the attachment is
    # the readable version: one table per module with the assumption, confidence and sources
    sheet = (PROGRAM / "kickoff-assumed-answers.md").read_text(encoding="utf-8").splitlines()
    sheet = sheet[sheet.index(next(l for l in sheet if l.startswith("# "))) + 1:]
    render_markdown_body(doc, sheet, NAVY)

    target = PROGRAM / "kickoff-pack.docx"
    doc.save(target)
    print(target)


if __name__ == "__main__":
    main()
