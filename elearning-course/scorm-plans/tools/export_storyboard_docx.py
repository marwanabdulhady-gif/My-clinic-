#!/usr/bin/env python3
"""Export all My Clinic Markdown storyboards into client-shareable Word review documents.

Requires python-docx and Pillow; see requirements-docx.txt. Example:
  python3 -m pip install -r elearning-course/scorm-plans/tools/requirements-docx.txt
  python3 elearning-course/scorm-plans/tools/export_storyboard_docx.py

The Markdown remains the editable source of truth. Generated .docx files are review artefacts.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
PLANS = ROOT / "elearning-course" / "scorm-plans"
OUTPUT = PLANS / "client-review"
LOGO = PLANS / "00_program" / "brand" / "logo-preview.png"
LOGO_HORIZONTAL = PLANS / "00_program" / "brand" / "assets" / "logo-horizontal-navy.png"
DATE = "14 October 2026"
REVISION = "v0.2"

MODULES = [
    ("01_health-and-safety-1", "01", "Facility_Management_and_Safety", "#008C99"),
    ("02_health-and-safety-2", "02", "Safety_in_Practice", "#008C99"),
    ("03_hr-employee-induction", "03", "Welcome_On_Board", "#652D8C"),
    ("04_operations-and-my-care", "04", "Operations_and_My_Care", "#9D3293"),
    ("05_rcm-documentation-and-coding", "05", "RCM_Part_A", "#7C0041"),
    ("06_rcm-ebp-benefits-and-rejections", "06", "RCM_Part_B", "#7C0041"),
    ("07_quality-safety-and-risk", "07", "Clinical_Governance", "#652D8C"),
    ("08_it-ehr-consultation", "08", "EHR_Consultation_Workflow", "#006E78"),
]


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.lstrip("#"))


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[[^]]*\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.replace("\\|", "|").strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str) -> None:
    """Insert a Word field (PAGE / NUMPAGES) so the footer shows live page numbers."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def brand_header(section, module_number: str | None, doc_id: str) -> None:
    """Running header: the supplied client logo on every page, document id on the right."""
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    usable = section.page_width - section.left_margin - section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    if LOGO_HORIZONTAL.exists():
        paragraph.add_run().add_picture(str(LOGO_HORIZONTAL), width=Inches(1.35))
    run = paragraph.add_run("\tMY CLINIC  |  STORYBOARD REVIEW  |  " + doc_id
                            + (f"  |  MODULE {module_number}" if module_number else ""))
    run.font.name = "Montserrat"
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb("5B6B7B")


def brand_footer(section, note: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(note + "   ·   Page ")
    run.font.name = "Montserrat"
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb("5B6B7B")
    add_field(paragraph, " PAGE ")
    mid = paragraph.add_run(" of ")
    mid.font.name = "Montserrat"
    mid.font.size = Pt(7.5)
    mid.font.color.rgb = rgb("5B6B7B")
    add_field(paragraph, " NUMPAGES ")
    for item in paragraph.runs:
        item.font.name = "Montserrat"
        item.font.size = Pt(7.5)
        item.font.color.rgb = rgb("5B6B7B")


def doc_control_table(doc, rows: list[tuple[str, str]], colour: str) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], colour.lstrip("#"))
        for run in cells[0].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        for para in cells[1].paragraphs:
            for run in para.runs:
                run.font.size = Pt(8.5)


def make_doc(module_number: str, module_slug: str, colour: str, source: Path, target: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Montserrat"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Montserrat")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11), ("Heading 4", 10)):
        style = doc.styles[name]
        style.font.name = "Montserrat"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Montserrat")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb("003867")
    if "Storyboard Code" not in doc.styles:
        code = doc.styles.add_style("Storyboard Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(7.5)
        code.paragraph_format.space_after = Pt(2)

    doc_id = f"MC-SB-{module_number} {REVISION}"
    brand_header(section, module_number, doc_id)
    brand_footer(section, f"Draft for SME review · {DATE} · {doc_id} · Module {module_number}")

    band = doc.add_table(rows=1, cols=1)
    set_cell_shading(band.rows[0].cells[0], colour.lstrip("#"))
    band.rows[0].cells[0].paragraphs[0].add_run("")
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.add_run().add_picture(str(LOGO), width=Inches(1.72))
    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(3)
    title.add_run(f"My Clinic — Module {module_number} storyboard")
    strap = doc.add_paragraph()
    strap.add_run("Client review document · ").bold = True
    strap.add_run("Generated from the controlled Markdown storyboard. The Markdown and JSON remain the editable source of truth. The supplied My Clinic logo appears in the running header of every page and is never regenerated or redrawn.")
    strap.paragraph_format.space_after = Pt(10)
    doc_control_table(doc, [
        ("Document", f"MC-SB-{module_number} · storyboard review rendering"),
        ("Revision", f"{REVISION} · {DATE}"),
        ("Status", "Draft for SME review — not approved course content"),
        ("Source of truth", f"{module_number}_*/storyboard/storyboard.md (Markdown + JSON)"),
        ("Companion files", "screen-visuals.html · visual-plan.json · narration-script.json · asset-register.md · lint-report.json"),
        ("Review route", "Named SMEs per the module header, then L&D authoring; open source-lock items stay ⚠ PENDING"),
    ], colour)
    doc.add_page_break()

    render_markdown_body(doc, source.read_text(encoding="utf-8").splitlines(), colour)
    doc.save(target)


def render_markdown_body(doc, content: list[str], colour: str) -> None:
    """Render controlled Markdown (tables, headings, quotes, code fences) into the document body."""
    i = 0
    in_code = False
    table_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        width = max(len(row) for row in table_buffer)
        # Markdown separator row is always second; ignore it.
        data = [row + [""] * (width - len(row)) for row in table_buffer if not all(re.fullmatch(r"[-: ]+", x) for x in row)]
        if not data:
            table_buffer = []
            return
        table = doc.add_table(rows=0, cols=width)
        table.style = "Light Shading Accent 1"
        for r_index, row in enumerate(data):
            cells = table.add_row().cells
            for c_index, raw in enumerate(row):
                cells[c_index].text = clean_inline(raw)
                for para in cells[c_index].paragraphs:
                    para.paragraph_format.space_after = Pt(1)
                    for run in para.runs:
                        run.font.name = "Montserrat"
                        run.font.size = Pt(7.5)
                if r_index == 0:
                    set_cell_shading(cells[c_index], colour.lstrip("#"))
                    for run in cells[c_index].paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if r_index == 0:
                set_repeat_table_header(table.rows[0])
        table_buffer = []

    while i < len(content):
        line = content[i]
        # Tables are accumulated before any other Markdown interpretation.
        if line.startswith("|") and line.rstrip().endswith("|"):
            table_buffer.append([part.strip() for part in line.strip().strip("|").split("|")])
            i += 1
            continue
        flush_table()
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Storyboard Code")
            p.add_run(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("!["):
            # The actual client logo has already been embedded at the start.
            i += 1
            continue
        match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if match:
            level = len(match.group(1))
            heading = doc.add_heading(clean_inline(match.group(2)), level=min(level, 4))
            if level == 3:
                heading.paragraph_format.page_break_before = True
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(clean_inline(line[2:]))
            run.italic = True
            run.font.color.rgb = rgb("5B6B7B")
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(line[2:]))
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(re.sub(r"^\d+\. ", "", line)))
            i += 1
            continue
        if line.startswith("---"):
            doc.add_paragraph("────────────────────────────────────────────────────", style="Normal")
            i += 1
            continue
        p = doc.add_paragraph()
        p.add_run(clean_inline(line))
        i += 1
    flush_table()


def make_program_pack(target: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    doc.styles["Normal"].font.name = "Montserrat"
    doc.styles["Normal"].font.size = Pt(10)
    pack_id = f"MC-SB-00 {REVISION}"
    brand_header(sec, None, pack_id)
    brand_footer(sec, f"Client review edition · {DATE} · {pack_id}")
    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Inches(2.05))
    doc.add_heading("My Clinic e-learning storyboard review pack", 0)
    doc.add_paragraph(f"Client review edition · {DATE} · {pack_id}")
    doc_control_table(doc, [
        ("Document", "MC-SB-00 · programme review pack index"),
        ("Revision", f"{REVISION} · {DATE}"),
        ("Covers", "Modules 01–08 storyboard renderings plus this index"),
        ("Status", "Draft for SME review — not approved course content"),
    ], "003867")
    doc.add_paragraph("This pack indexes the eight controlled storyboard documents. It makes clear what is ready for SME review and what remains blocked by policy, rights, voice or test-environment decisions.")
    doc.add_heading("SME review queue", 1)
    doc.add_paragraph("Use 00_program/storyboard-review-register.md and its CSV sibling alongside this pack. The register routes screen-level interaction detail, visual-hotspot and privacy checks to named review roles; it is a pre-build queue, not an approval record.")
    doc.add_heading("Review documents", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    headers = ["Module", "Storyboard", "Review status", "Core hand-off"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        set_cell_shading(cell, "003867")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for folder, number, slug, _colour in MODULES:
        row = table.add_row().cells
        row[0].text = f"{number}"
        row[1].text = slug.replace("_", " ")
        row[2].text = "Draft — SME/source-lock review"
        row[3].text = "Screen cards · narration script · visual plan · asset register"
    doc.add_heading("Non-negotiable release gates", 1)
    for item in [
        "Do not convert a ⚠ PENDING source fact into approved course content. Each module source-lock record names its owner.",
        "Use the supplied My Clinic logo PNGs; never generate or redraw the logo.",
        "Approve a fixed Noor and Dr. Omar voice route before shipping audio. Draft scripts are ready; audio is not yet a release asset.",
        "Module 08 uses only test-environment screenshots with fictional test data and DPO review. The source video is not a production asset.",
        "Build and LMS delivery are Phase 4–5 work after SME approval, edumints lint, package QA and sign-off.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Files in this folder", 1)
    doc.add_paragraph("Each Module file is a client-shareable Word rendering of its complete Markdown storyboard. The corresponding source and screen visual review board live in the module’s storyboard folder.")
    doc.save(target)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    exported = []
    for folder, number, slug, colour in MODULES:
        source = PLANS / folder / "storyboard" / "storyboard.md"
        if not source.exists():
            raise SystemExit(f"Missing storyboard source: {source}")
        target = OUTPUT / f"MyClinic_{number}_{slug}_Storyboard_v0.2.docx"
        make_doc(number, slug, colour, source, target)
        exported.append(target.name)
    overview = OUTPUT / "MyClinic_Elearning_Storyboard_Review_Pack_v0.2.docx"
    make_program_pack(overview)
    readme = ["# My Clinic — client-review documents", "", "These Word documents are client-shareable renderings of the controlled Markdown storyboards. They were exported on 2026-10-14. The Markdown/JSON/HTML files in each module’s `storyboard/` directory remain the editable production source.", "", "Use the [SME/build-detail review register](../00_program/storyboard-review-register.md) and its [tracker-import CSV](../00_program/storyboard-review-register.csv) alongside the Word documents. Neither file records policy approval.", "", "To refresh the Word renderings after an approved source change, install `tools/requirements-docx.txt` and run `python3 elearning-course/scorm-plans/tools/export_storyboard_docx.py` from the repository root.", "", "| File | Contents |", "|---|---|"]
    readme.append(f"| [{overview.name}]({overview.name}) | Programme overview, review route and release gates |")
    for filename in exported:
        readme.append(f"| [{filename}]({filename}) | Complete module screen cards, visual specification, narration and hand-off notes |")
    readme += ["", "All documents remain **drafts for SME review** until their source-lock decisions, asset approvals and voice route are signed off.", ""]
    (OUTPUT / "README.md").write_text("\n".join(readme), encoding="utf-8")
    print("\n".join(str(OUTPUT / x) for x in [overview.name, *exported]))


if __name__ == "__main__":
    main()
